import asyncio
import json
from io import BytesIO
from unittest.mock import AsyncMock

import httpx
import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from sqlalchemy import func, select
from sqlalchemy.orm import Session, sessionmaker
from test_chat import mock_chat, provider_response

from app.ai.contracts import ModelResult, RAGAnswer
from app.core.exceptions import AppError
from app.models.ai import Document, DocumentChunk
from app.rag.embeddings import DIMENSIONS, Embeddings
from app.rag.parsing import MAX_BYTES, Chunk, chunk_pages, parse_document
from app.rag.repository import DocumentRepository
from app.schemas.documents import DocumentSource
from app.services.chat import ChatService


def test_txt_chunking_preserves_pages_and_overlap():
    pages = parse_document("policy.txt", ("\ufeff  Правила\n\nвозврата " * 200).encode())
    chunks = chunk_pages(pages)
    assert len(chunks) > 1
    assert all(len(c.content) <= 1200 for c in chunks)
    assert all(c.page_number is None for c in chunks)
    assert chunks[0].content[-150:] in chunks[1].content
    assert [c.page_number for c in chunk_pages([(1, "one"), (2, "two")])] == [1, 2]


@pytest.mark.parametrize(
    "filename,data",
    [
        ("../secret.txt", b"text"),
        ("a\\b.txt", b"text"),
        ("x.exe", b"text"),
        ("empty.txt", b"  \n"),
        ("bad.txt", b"\xff"),
        ("fake.pdf", b"text"),
        ("fake.docx", b"PKbad"),
        ("huge.txt", b"a" * (MAX_BYTES + 1)),
        ("long.txt", b"a" * 200001),
    ],
)
def test_rejects_invalid_documents(filename, data):
    with pytest.raises(AppError):
        parse_document(filename, data)


def test_docx_includes_tables_and_pdf_requires_text():
    doc = DocxDocument()
    doc.add_paragraph("Refund policy")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "10 working days"
    stream = BytesIO()
    doc.save(stream)
    assert "10 working days" in parse_document("rules.docx", stream.getvalue())[0][1]
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(AppError, match="No readable text"):
        parse_document("scan.pdf", stream.getvalue())


def test_store_is_atomic_and_duplicate_upload_is_idempotent(database):
    repo = DocumentRepository(sessionmaker(database))
    chunks = [Chunk("Refund in 10 days", 1)]
    vector = [1.0] + [0.0] * (DIMENSIONS - 1)
    first = repo.store("rules.txt", b"same", chunks, [vector])
    assert repo.store("renamed.txt", b"same", chunks, [vector]).id == first.id
    with pytest.raises(ValueError):
        repo.store("bad.txt", b"different", chunks, [])
    with Session(database) as session:
        assert session.scalar(select(func.count()).select_from(Document)) == 1
        assert session.scalar(select(func.count()).select_from(DocumentChunk)) == 1
    if database.dialect.name == "postgresql":
        assert repo.search(vector, 5, 0.5)[0].document_id == first.id
        assert repo.search([-x for x in vector], 5, 0.5) == []


def test_embedding_api_and_validation(monkeypatch):
    def handler(request):
        body = json.loads(request.content)
        assert body["model"] == "text-embedding-3-small"
        assert body["dimensions"] == DIMENSIONS
        return httpx.Response(
            200,
            json={
                "data": [
                    {
                        "index": 0,
                        "embedding": [1.0] + [0.0] * (DIMENSIONS - 1),
                        "object": "embedding",
                    }
                ],
                "model": body["model"],
                "object": "list",
                "usage": {"prompt_tokens": 3, "total_tokens": 3},
            },
        )

    with mock_chat(monkeypatch, handler) as client:
        embeddings = client.app.state.document_service.embeddings
        assert len(asyncio.run(embeddings.embed(["policy"]))[0]) == DIMENSIONS
    with pytest.raises(AppError, match="not configured"):
        asyncio.run(Embeddings(None).embed(["policy"]))


@pytest.mark.parametrize("code,expected", [(429, 429), (500, 503)])
def test_embedding_errors(monkeypatch, code, expected):
    with mock_chat(
        monkeypatch, lambda _: httpx.Response(code, json={"error": {"message": "private"}})
    ) as client:
        response = client.post("/api/v1/documents/search", json={"query": "policy"})
        assert response.status_code == expected
        assert "private" not in response.text


def source():
    return DocumentSource(
        document_id=1,
        chunk_id=3,
        filename="refund.txt",
        page_number=None,
        content="Refund in 10 working days",
        similarity=0.8,
    )


@pytest.mark.parametrize(
    "ids,supported,status",
    [
        ([3], True, "grounded_answer"),
        ([], True, "insufficient_data"),
        ([3], False, "insufficient_data"),
    ],
)
def test_rag_validates_sources(ids, supported, status):
    model = AsyncMock()
    model.generate.return_value = ModelResult(
        RAGAnswer(answer="10 working days", supported=supported, chunk_ids=ids),
        False,
        "test",
        "test",
        1,
        1,
    )
    documents = AsyncMock()
    documents.search.return_value = [source()]
    response = asyncio.run(ChatService(model, documents).reply_with_documents("When?"))
    assert response.status == status
    assert len(response.sources) == (1 if status == "grounded_answer" else 0)


def test_unknown_citations_and_empty_retrieval():
    model = AsyncMock()
    documents = AsyncMock()
    documents.search.return_value = []
    service = ChatService(model, documents)
    assert asyncio.run(service.reply_with_documents("When?")).status == "insufficient_data"
    model.generate.assert_not_called()
    documents.search.return_value = [source()]
    model.generate.return_value = ModelResult(
        RAGAnswer(answer="invented", supported=True, chunk_ids=[999]), False, "test", "test", 1, 1
    )
    with pytest.raises(AppError, match="Unknown source"):
        asyncio.run(service.reply_with_documents("When?"))


def test_rag_chat_uses_real_structured_output_adapter(monkeypatch):
    def handler(request):
        body = json.loads(request.content)
        assert "supported" in body["text"]["format"]["schema"]["required"]
        assert "Refund in 10 working days" in body["input"][0]["content"]
        return httpx.Response(
            200,
            json=provider_response(
                {"answer": "10 working days", "supported": True, "chunk_ids": [3]}
            ),
        )

    with mock_chat(monkeypatch, handler) as client:
        client.app.state.document_service.search = AsyncMock(return_value=[source()])
        response = client.post("/api/v1/chat", json={"message": "When?", "use_documents": True})
        assert response.status_code == 200
        assert response.json()["sources"][0]["chunk_id"] == 3


def test_upload_rejects_bad_inputs_before_provider(monkeypatch):
    def handler(_):
        pytest.fail("Invalid uploads must not reach OpenAI")

    with mock_chat(monkeypatch, handler) as client:
        assert client.post("/api/v1/documents?filename=../x.txt", content=b"x").status_code == 422
        assert (
            client.post(
                "/api/v1/documents?filename=x.txt", content=b"x" * (MAX_BYTES + 1)
            ).status_code
            == 413
        )
        assert client.post("/api/v1/documents?filename=x.pdf", content=b"fake").status_code == 422
        assert (
            client.post("/api/v1/documents/search", json={"query": "x", "top_k": 100}).status_code
            == 422
        )


def test_synthetic_pdfs_are_searchable_text():
    from pathlib import Path

    paths = list((Path(__file__).resolve().parents[2] / "documents/synthetic").glob("*.pdf"))
    assert len(paths) == 5
    for path in paths:
        pages = parse_document(path.name, path.read_bytes())
        assert pages[0][0] == 1
        assert "synthetic demonstration" in pages[0][1]
        assert len(pages[0][1]) > 400


def test_upload_pipeline_persists_and_lists_documents(monkeypatch, database):
    def handler(request):
        body = json.loads(request.content)
        return httpx.Response(
            200,
            json={
                "object": "list",
                "model": body["model"],
                "data": [
                    {
                        "object": "embedding",
                        "index": i,
                        "embedding": [1.0] + [0.0] * (DIMENSIONS - 1),
                    }
                    for i in range(len(body["input"]))
                ],
                "usage": {"prompt_tokens": 5, "total_tokens": 5},
            },
        )

    with mock_chat(monkeypatch, handler) as client:
        client.app.state.document_service.repository = DocumentRepository(sessionmaker(database))
        response = client.post(
            "/api/v1/documents?filename=policy.txt", content=b"Refund within 10 days"
        )
        assert response.status_code == 201
        assert response.json()["status"] == "ready"
        listed = client.get("/api/v1/documents").json()
        assert listed == [response.json()]
        assert client.get("/api/v1/documents?offset=1").json() == []


@pytest.mark.parametrize(
    "vector,index", [([1.0], 0), ([0.0] * DIMENSIONS, 0), ([1.0] * DIMENSIONS, 7)]
)
def test_malformed_embedding_response(monkeypatch, vector, index):
    def handler(_):
        return httpx.Response(
            200,
            json={
                "object": "list",
                "model": "text-embedding-3-small",
                "data": [{"object": "embedding", "index": index, "embedding": vector}],
                "usage": {"prompt_tokens": 1, "total_tokens": 1},
            },
        )

    with mock_chat(monkeypatch, handler) as client:
        response = client.post("/api/v1/documents/search", json={"query": "refund"})
        assert response.status_code == 502
        assert response.json()["error"]["code"] == "ai_invalid_response"
