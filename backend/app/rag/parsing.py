"""Bounded, in-memory parsing; uploaded filenames never become filesystem paths."""

import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.exceptions import AppError

MAX_BYTES = 5 * 1024 * 1024
MAX_TEXT = 200_000
CONTENT_TYPES = {
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def invalid(message: str, status: int = 422) -> AppError:
    return AppError(status_code=status, code="invalid_document", message=message)


def validate_filename(filename: str) -> str:
    if not filename or len(filename) > 250 or any(c in filename for c in "/\\\x00"):
        raise invalid("Provide a filename without path components")
    if any(ord(c) < 32 for c in filename):
        raise invalid("Invalid filename")
    extension = PurePosixPath(filename).suffix.lower()
    if extension not in CONTENT_TYPES:
        raise invalid("Supported formats: UTF-8 TXT, PDF, DOCX", 415)
    return extension


@dataclass(frozen=True)
class Chunk:
    content: str
    page_number: int | None


def parse_document(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    extension = validate_filename(filename)
    if len(data) > MAX_BYTES:
        raise invalid("Document exceeds 5 MiB", 413)
    if not data:
        raise invalid("Document is empty")
    try:
        if extension == ".txt":
            pages = [(None, data.decode("utf-8-sig"))]
        elif extension == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise invalid("Invalid PDF signature")
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted or len(reader.pages) > 100:
                raise invalid("Encrypted PDFs or PDFs over 100 pages are unsupported")
            pages = []
            total = 0
            for number, page in enumerate(reader.pages, 1):
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > 10 * MAX_BYTES:
                    raise invalid("PDF page is too complex")
                text = page.extract_text() or ""
                total += len(text)
                if total > MAX_TEXT:
                    raise invalid("Extracted text exceeds 200000 characters")
                pages.append((number, text))
        else:
            with ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(e.file_size for e in entries) > 20 * MAX_BYTES:
                    raise invalid("DOCX archive exceeds extraction limits")
                if "word/document.xml" not in archive.namelist():
                    raise invalid("Invalid DOCX structure")
            document = DocxDocument(BytesIO(data))
            text = "\n".join(p.text for p in document.paragraphs)
            text += "\n" + "\n".join(
                " | ".join(c.text for c in row.cells)
                for table in document.tables
                for row in table.rows
            )
            pages = [(None, text)]
    except AppError:
        raise
    except Exception:  # noqa: BLE001 - third-party parser input boundary
        raise invalid("Cannot parse document; provide a valid, unencrypted text document") from None
    if sum(len(text) for _, text in pages) > MAX_TEXT:
        raise invalid("Extracted text exceeds 200000 characters")
    cleaned = [
        (page, re.sub(r"\s+", " ", text.replace("\x00", " ")).strip()) for page, text in pages
    ]
    cleaned = [(page, text) for page, text in cleaned if text]
    if not cleaned:
        raise invalid("No readable text; scanned PDFs require OCR, which is not supported")
    return cleaned


def chunk_pages(
    pages: list[tuple[int | None, str]], size: int = 1200, overlap: int = 200
) -> list[Chunk]:
    if not 0 <= overlap < size:
        raise ValueError("Require 0 <= overlap < size")
    chunks = []
    for page, text in pages:
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            if end < len(text):
                boundary = text.rfind(" ", start + size // 2, end)
                if boundary > start:
                    end = boundary
            content = text[start:end].strip()
            if content:
                chunks.append(Chunk(content, page))
            if end == len(text):
                break
            start = max(start + 1, end - overlap)
    return chunks
